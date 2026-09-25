from pathlib import Path
from PIL import Image, ImageOps
from openpyxl import load_workbook
from difflib import SequenceMatcher
import json
import re
import shutil
import unicodedata
from denomination_assets import prepare_denomination_logos, match_denomination_logo, normalize as normalize_denomination


SOURCE = Path("/Users/joshuagbafa/Downloads/Bishop's project/MASTER PORTRAITS IMPORT/PASTORS PICTURES")
WORKBOOK = Path("/Users/joshuagbafa/Downloads/PASTORS DATA.xlsx")
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "bishops"
if not SOURCE.exists() or any(Path('/Users/joshuagbafa/Downloads').glob('*.crdownload')):
    raise SystemExit('Finish and extract the master download before rebuilding portraits.')
OUT.mkdir(parents=True, exist_ok=True)
for old_image in OUT.glob("*.jpg"):
    old_image.unlink()


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def name_case(value):
    words = []
    for part in clean(value).split():
        words.append("-".join(piece.capitalize() for piece in part.split("-")))
    return " ".join(words)


def safe_age(value):
    text = clean(value)
    try:
        age = int(float(text))
        return str(age) if 18 <= age <= 100 else ""
    except (TypeError, ValueError):
        return ""


def normalize(value):
    value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode().upper()
    value = re.sub(r"\[[^]]+\]|\([^)]*\)|\.(JPG|JPEG|PNG|HEIF|HEIC)$", " ", value)
    value = re.sub(r"COPY OF|RED\s*JACKET|WHATSAPP IMAGE|PHOTO|PICTURES|PICTURE|PHOTOS", " ", value)
    value = re.sub(r"\b(BISHOP|BISHOPS|BS|REV|REVEREND|PASTOR|PS|SISTER|ES|DR|PROPHET|MOTHER|JNR|JR)\b", " ", value)
    value = re.sub(r"\b(IMG|DSC|MEDIA)\b|\d+", " ", value)
    return " ".join(re.findall(r"[A-Z]+", value))


def score(name, label):
    a, b = normalize(name), normalize(label)
    if not a or not b:
        return 0
    if a == b:
        return 1
    if a.replace(" ", "") == b.replace(" ", ""):
        return .99
    at, bt = set(a.split()), set(b.split())
    overlap = at & bt
    if len(overlap) >= 2 and (at <= bt or bt <= at):
        return .94
    token_score = 2 * len(overlap) / (len(at) + len(bt))
    sequence_score = SequenceMatcher(None, a, b).ratio()
    if len(overlap) >= 2:
        token_score += .08
    return min(.93, max(token_score, sequence_score))


def labels_for(path):
    labels = [path.stem]
    parent = path.parent
    while parent != SOURCE and len(labels) < 5:
        labels.append(parent.name)
        parent = parent.parent
    return labels


def usable_images():
    images = []
    for path in SOURCE.rglob("*"):
        if path.suffix.lower() not in {".jpg", ".jpeg", ".png", ".heif", ".heic"}:
            continue
        try:
            with Image.open(path) as image:
                image.verify()
        except Exception:
            continue
        images.append({"path": path, "labels": labels_for(path)})
    return images


workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
sheet = workbook.active
rows = iter(sheet.iter_rows(values_only=True))
headers = [clean(value) for value in next(rows)]
people = []
for values in rows:
    row = dict(zip(headers, values))
    is_leader = (
        "BISHOP" in clean(row.get("ADMINRANK")).upper()
        or clean(row.get("STATUSRANK")).upper() == "BISHOP"
        or bool(row.get("YEARCONSECRATED"))
    )
    if not is_leader or not row.get("FULLNAME"):
        continue
    denomination = clean(row.get("DENOMINATION"))
    source_group = "first_love" if "FIRST LOVE" in denomination.upper() else "ud"
    branch = re.sub(r"(?i)first\s*love", "Central Church", name_case(row.get("BRANCH")))
    people.append({
        "sourceId": clean(row.get("PASTORID")),
        "name": name_case(row.get("FULLNAME")),
        "denomination": denomination or "Leadership Network",
        "region": name_case(row.get("COUNTRY")) or "International",
        "branch": branch,
        "sourceStatus": clean(row.get("PASTORSTATUS")) or "ACTIVE",
        "yearConsecrated": clean(row.get("YEARCONSECRATED")),
        "gender": clean(row.get("GENDER")).upper(),
        "sourceGroup": source_group,
        "age": safe_age(row.get("AGE")),
        "mobile": clean(row.get("MOBILE")),
        "whatsapp": clean(row.get("WHATSAPP NUMBER")),
        "email": clean(row.get("ALTERNATEEMAIL")),
        "yearAppointed": clean(row.get("YEARAPPOINTED")),
        "yearOrdained": clean(row.get("YEARORDAINED")),
        "adminRank": clean(row.get("ADMINRANK")),
        "ministryRank": clean(row.get("MINISTRYRANK")),
        "statusRank": clean(row.get("STATUSRANK")),
        "functionsRank": clean(row.get("FUNCTIONSRANK")),
        "council": clean(row.get("COUNCIL")),
        "diocese": clean(row.get("DIOCESE")),
        "city": name_case(row.get("CITY")),
        "state": name_case(row.get("STATE")),
        "address": clean(row.get("ADDRESS")),
        "qualification": clean(row.get("QUALIFICATION")),
        "profession": clean(row.get("PROFESSION")),
        "occupation": clean(row.get("OCCUPATION")),
        "maritalStatus": clean(row.get("MARITALSTATUS")),
    })

photos = usable_images()
denomination_logos = prepare_denomination_logos(ROOT)
report = []
entries = []
for person in people:
    candidates = []
    for photo in photos:
        best = max((score(person["name"], label) for label in photo["labels"]), default=0)
        if best >= .58:
            candidates.append((best, photo))
    candidates.sort(key=lambda item: item[0], reverse=True)
    chosen = candidates[0] if candidates and candidates[0][0] >= .72 else None
    first_love_candidates = [
        item for item in candidates
        if "FIRST LOVE" in str(item[1]["path"]).upper()
    ]
    first_love_candidates.sort(key=lambda item: item[0], reverse=True)
    first_love_by_folder = bool(first_love_candidates and first_love_candidates[0][0] >= .72)
    if first_love_by_folder:
        chosen = first_love_candidates[0]
    image_path = None
    if chosen:
        source = chosen[1]["path"]
        try:
            with Image.open(source) as image:
                image = ImageOps.exif_transpose(image).convert("RGB")
                scale = max(720 / image.width, 900 / image.height)
                resized = image.resize((round(image.width * scale), round(image.height * scale)), Image.Resampling.LANCZOS)
                left = max(0, (resized.width - 720) // 2)
                top = max(0, min((resized.height - 900) // 5, resized.height - 900))
                crop = resized.crop((left, top, left + 720, top + 900))
                filename = f"{len(entries) + 1:03d}.jpg"
                crop.save(OUT / filename, "JPEG", quality=88, optimize=True, progressive=True)
                image_path = f"assets/bishops/{filename}"
        except Exception:
            chosen = None
    source_image = str(chosen[1]["path"].relative_to(SOURCE)) if chosen else ""
    first_love = first_love_by_folder or person["sourceGroup"] == "first_love"
    honorific = "Mother" if first_love and person["gender"] == "FEMALE" else ""
    designation = "UD-UO-FLC190" if first_love else "UD-OLGCA"
    organization = designation
    code = len(entries) + 1
    entries.append({
        "code": code,
        "name": person["name"],
        "organization": organization,
        "denomination": person["denomination"],
        "denominationLogo": denomination_logos.get(normalize_denomination('FIRST LOVE CHURCH WORLDWIDE')) if first_love else match_denomination_logo(person["denomination"], denomination_logos),
        "region": person["region"],
        "branch": person["branch"],
        "image": image_path,
        "amount": 100,
        "sourceStatus": person["sourceStatus"],
        "yearConsecrated": person["yearConsecrated"],
        "honorific": honorific,
        "designation": designation,
        "gender": person["gender"],
        "age": person["age"],
        "mobile": person["mobile"],
        "whatsapp": person["whatsapp"],
        "email": person["email"],
        "yearAppointed": person["yearAppointed"],
        "yearOrdained": person["yearOrdained"],
        "adminRank": person["adminRank"],
        "ministryRank": person["ministryRank"],
        "statusRank": person["statusRank"],
        "functionsRank": person["functionsRank"],
        "council": person["council"],
        "diocese": person["diocese"],
        "city": person["city"],
        "state": person["state"],
        "address": person["address"],
        "qualification": person["qualification"],
        "profession": person["profession"],
        "occupation": person["occupation"],
        "maritalStatus": person["maritalStatus"],
    })
    report.append({
        "code": code,
        "sourceId": person["sourceId"],
        "name": person["name"],
        "matched": bool(image_path),
        "score": round(chosen[0], 3) if chosen else None,
        "sourceImage": source_image or None,
        "honorific": honorific,
        "designation": designation,
        "otherCandidates": [
            {"score": round(item[0], 3), "path": str(item[1]["path"].relative_to(SOURCE))}
            for item in candidates[1:4]
        ],
    })

(ROOT / "data" / "bishops.js").write_text(
    "window.BISHOPS = " + json.dumps(entries, ensure_ascii=False, separators=(",", ":")) + ";\n"
)
(ROOT / "data" / "leader-photo-match-report.json").write_text(json.dumps(report, indent=2))

logo_source = Path("/Users/joshuagbafa/Downloads/Bishop's project/Brand/IMG_6513.jpg")
shutil.copy2(logo_source, ROOT / "assets" / "mitre.jpg")

matched = sum(1 for item in entries if item["image"])
print(json.dumps({"leaders": len(entries), "matchedPortraits": matched, "needsPortraitReview": len(entries) - matched, "availableImages": len(photos)}, indent=2))
