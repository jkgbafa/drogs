from pathlib import Path
from openpyxl import load_workbook
from PIL import Image, ImageOps
from difflib import SequenceMatcher
import json
import re
import unicodedata
import argparse
from denomination_assets import prepare_denomination_logos, match_denomination_logo, normalize as normalize_denomination


SOURCE = Path("/Users/joshuagbafa/Downloads/PASTORS DATA.xlsx")
FIRST_LOVE_SOURCE = Path("/Users/joshuagbafa/Downloads/FIRST LOVE BISHOPS AND PASTORS.xlsx")
PHOTO_SOURCE = Path("/Users/joshuagbafa/Downloads/Bishop's project/MASTER PORTRAITS IMPORT/PASTORS PICTURES")
ROOT = Path(__file__).resolve().parents[1]
PHOTO_OUT = ROOT / "assets" / "pastors"
parser = argparse.ArgumentParser()
parser.add_argument('--roster-only', action='store_true', help='Update the roster while the master portraits are downloading.')
args = parser.parse_args()


def clean(value):
    if value is None:
        return ""
    value = re.sub(r"\s+", " ", str(value)).strip()
    value = re.sub(r"\s+\.\s+", " ", value)
    return value


def name_case(value):
    words = []
    for part in clean(value).split():
        words.append("-".join(piece.capitalize() for piece in part.split("-")))
    return " ".join(words)


def safe_age(value):
    text = clean(value)
    try:
        age = int(float(text))
        return str(age) if 18 <= age <= 100 else ""
    except (TypeError, ValueError):
        return ""


def normalize_name(value):
    value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode().upper()
    value = re.sub(r"\.(JPG|JPEG|PNG|WEBP|HEIC|HEIF)$", "", value)
    value = re.sub(r"\b(LADY|PASTOR|BISHOP|REVEREND|REV|PS|LP|LR|MS|MINISTER|SHEPHERD|SISTER|DR|PROPHET|EVANGELIST)\b", " ", value)
    value = re.sub(r"\[[^]]+\]|\([^)]*\)|\d+", " ", value)
    return " ".join(re.findall(r"[A-Z]+", value))


def photo_labels(path):
    stem = path.stem
    labels = [stem]
    if " - " in stem:
        labels.append(stem.split(" - ", 1)[0])
    return [normalize_name(label) for label in labels if normalize_name(label)]


def match_score(name, labels):
    target = normalize_name(name)
    target_tokens = set(target.split())
    best = 0
    for label in labels:
        if target == label:
            return 1
        if target.replace(" ", "") == label.replace(" ", ""):
            return .99
        label_tokens = set(label.split())
        overlap = target_tokens & label_tokens
        if len(overlap) >= 2 and (target_tokens <= label_tokens or label_tokens <= target_tokens):
            best = max(best, .96)
        token_score = 2 * len(overlap) / max(1, len(target_tokens) + len(label_tokens))
        sequence_score = SequenceMatcher(None, target, label).ratio()
        if len(overlap) >= 2:
            token_score += .08
        best = max(best, min(.95, token_score), sequence_score)
    return best


def usable_photos():
    photos = []
    if not PHOTO_SOURCE.exists():
        return photos
    for path in PHOTO_SOURCE.rglob("*"):
        if not path.is_file() or path.name.startswith("."):
            continue
        try:
            with Image.open(path) as image:
                image.verify()
        except Exception:
            continue
        labels = photo_labels(path)
        if labels:
            photos.append({"path": path, "labels": labels})
    return photos


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
headers = [clean(cell.value) for cell in sheet[1]]
positions = {header: index for index, header in enumerate(headers)}

entries = []
seen = set()
denomination_logos = prepare_denomination_logos(ROOT)
first_love_logo = denomination_logos.get(normalize_denomination("FIRST LOVE CHURCH WORLDWIDE"))
for row in sheet.iter_rows(min_row=2, values_only=True):
    row = tuple(row) + (None,) * max(0, len(headers) - len(row))
    is_leader = (
        "BISHOP" in clean(row[positions["ADMINRANK"]]).upper()
        or clean(row[positions["STATUSRANK"]]).upper() == "BISHOP"
        or bool(row[positions["YEARCONSECRATED"]])
    )
    if is_leader:
        continue
    name = name_case(row[positions["FULLNAME"]])
    denomination = clean(row[positions["DENOMINATION"]])
    branch = name_case(row[positions["BRANCH"]])
    branch = re.sub(r"(?i)first\s*love", "Central Church", branch)
    country = name_case(row[positions["COUNTRY"]])
    if not name or not denomination:
        continue
    key = (re.sub(r"[^a-z]", "", name.lower()), denomination.lower(), branch.lower(), country.lower())
    if key in seen:
        continue
    seen.add(key)
    entries.append({
        "code": len(entries) + 1,
        "name": name,
        "organization": "UD-OLGCA",
        "designation": "UD-OLGCA",
        "denomination": denomination,
        "denominationLogo": match_denomination_logo(denomination, denomination_logos),
        "region": country or "International",
        "branch": branch,
        "image": None,
        "amount": 50,
        "yearAppointed": clean(row[positions["YEARAPPOINTED"]]),
        "yearOrdained": clean(row[positions["YEARORDAINED"]]),
        "gender": clean(row[positions["GENDER"]]).upper(),
        "age": safe_age(row[positions["AGE"]]),
        "mobile": clean(row[positions["MOBILE"]]),
        "whatsapp": clean(row[positions["WHATSAPP NUMBER"]]),
        "email": clean(row[positions["ALTERNATEEMAIL"]]),
        "adminRank": clean(row[positions["ADMINRANK"]]),
        "ministryRank": clean(row[positions["MINISTRYRANK"]]),
        "statusRank": clean(row[positions["STATUSRANK"]]),
        "functionsRank": clean(row[positions["FUNCTIONSRANK"]]),
        "council": clean(row[positions["COUNCIL"]]),
        "diocese": clean(row[positions["DIOCESE"]]),
        "city": name_case(row[positions["CITY"]]),
        "state": name_case(row[positions["STATE"]]),
        "address": clean(row[positions["ADDRESS"]]),
        "qualification": clean(row[positions["QUALIFICATION"]]),
        "profession": clean(row[positions["PROFESSION"]]),
        "occupation": clean(row[positions["OCCUPATION"]]),
        "maritalStatus": clean(row[positions["MARITALSTATUS"]]),
        "supervisingBishop": "",
    })

# The First Love sheet is the authoritative organization roster. Existing master
# records retain their richer personal data; people absent from the master are
# added with the country and supervising bishop supplied by First Love.
if FIRST_LOVE_SOURCE.exists():
    first_love_sheet = load_workbook(FIRST_LOVE_SOURCE, read_only=True, data_only=True).active
    first_love_sheet.reset_dimensions()
    by_name = {}
    for entry in entries:
        by_name.setdefault(normalize_name(entry["name"]), []).append(entry)
    supervising_bishop = ""
    for values in first_love_sheet.iter_rows(min_row=2, values_only=True):
        values = tuple(values) + (None,) * max(0, 4 - len(values))
        if clean(values[0]):
            supervising_bishop = name_case(values[0])
        name = name_case(values[2])
        country = name_case(values[1])
        if not name or name.upper().startswith("TOTAL NO."):
            continue
        matches = by_name.get(normalize_name(name), [])
        entry = next((item for item in matches if normalize_name(item["region"]) == normalize_name(country)), matches[0] if matches else None)
        if entry is None:
            entry = {
                "code": 0,
                "name": name,
                "organization": "UD-UO-FLC190",
                "designation": "UD-UO-FLC190",
                "denomination": "FIRST LOVE CHURCH",
                "denominationLogo": first_love_logo,
                "region": country or "International",
                "branch": "",
                "image": None,
                "amount": 50,
                "yearAppointed": "", "yearOrdained": "", "gender": "", "age": "",
                "mobile": "", "whatsapp": "", "email": "", "adminRank": "",
                "ministryRank": "", "statusRank": "PASTOR", "functionsRank": "",
                "council": "", "diocese": "", "city": "", "state": "", "address": "",
                "qualification": "", "profession": "", "occupation": "", "maritalStatus": "",
                "supervisingBishop": supervising_bishop,
            }
            entries.append(entry)
            by_name.setdefault(normalize_name(name), []).append(entry)
        else:
            entry.update({
                "organization": "UD-UO-FLC190",
                "designation": "UD-UO-FLC190",
                "denomination": "FIRST LOVE CHURCH",
                "denominationLogo": first_love_logo,
                "supervisingBishop": supervising_bishop,
            })
            if not entry["region"] and country:
                entry["region"] = country

for code, entry in enumerate(entries, 1):
    entry["code"] = code

if args.roster_only:
    (ROOT / "data" / "pastors.js").write_text(
        "window.PASTORS = " + json.dumps(entries, ensure_ascii=False, separators=(",", ":")) + ";\n"
    )
    print(json.dumps({"pastors": len(entries), "firstLove": sum(e['organization'] == 'UD-UO-FLC190' for e in entries), "portraits": "Pending complete master download"}))
    raise SystemExit(0)

if not PHOTO_SOURCE.exists() or any(Path('/Users/joshuagbafa/Downloads').glob('*.crdownload')):
    raise SystemExit('Finish and extract the master download before rebuilding portraits; use --roster-only for roster updates.')
PHOTO_OUT.mkdir(parents=True, exist_ok=True)
for old in PHOTO_OUT.glob("*"):
    if old.is_file():
        old.unlink()

photos = usable_photos()
token_index = {}
for photo_index, photo in enumerate(photos):
    for token in set(" ".join(photo["labels"]).split()):
        if len(token) >= 3:
            token_index.setdefault(token, set()).add(photo_index)

photo_report = []
for entry in entries:
    tokens = [token for token in normalize_name(entry["name"]).split() if len(token) >= 3]
    candidate_indices = set()
    for token in tokens:
        candidate_indices.update(token_index.get(token, set()))
    ranked = sorted(
        ((match_score(entry["name"], photos[index]["labels"]), index) for index in candidate_indices),
        reverse=True,
    )
    chosen = ranked[0] if ranked and ranked[0][0] >= .80 else None
    source_path = None
    if chosen:
        source_path = photos[chosen[1]]["path"]
        try:
            with Image.open(source_path) as image:
                image = ImageOps.exif_transpose(image).convert("RGB")
                scale = max(480 / image.width, 600 / image.height)
                resized = image.resize((round(image.width * scale), round(image.height * scale)), Image.Resampling.LANCZOS)
                left = max(0, (resized.width - 480) // 2)
                top = max(0, min((resized.height - 600) // 5, resized.height - 600))
                crop = resized.crop((left, top, left + 480, top + 600))
                filename = f'{entry["code"]:04d}.jpg'
                crop.save(PHOTO_OUT / filename, "JPEG", quality=75, optimize=True, progressive=True)
                entry["image"] = f"assets/pastors/{filename}"
        except Exception:
            chosen = None
            source_path = None
    photo_report.append({
        "code": entry["code"],
        "name": entry["name"],
        "matched": bool(entry["image"]),
        "score": round(chosen[0], 3) if chosen else None,
        "sourceImage": str(source_path.relative_to(PHOTO_SOURCE)) if source_path else None,
    })

(ROOT / "data" / "pastor-photo-match-report.json").write_text(json.dumps(photo_report, indent=2))
(ROOT / "data" / "pastors.js").write_text(
    "window.PASTORS = " + json.dumps(entries, ensure_ascii=False, separators=(",", ":")) + ";\n"
)
matched = sum(1 for entry in entries if entry["image"])
print(json.dumps({"pastors": len(entries), "matchedPortraits": matched, "needsPortraitReview": len(entries) - matched, "availableImages": len(photos), "first": entries[:2], "last": entries[-1:]}, indent=2))
